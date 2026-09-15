"""单元评估快照与 SEC 单元报告的测试（走真实库：需先 init + train）。"""
from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from src.api import services as S          # noqa: E402
from src.report import unit_report as R    # noqa: E402


def _key():
    unit = S._units()[0]["unit_id"].iloc[0]
    as_of, deck = S._as_of(S._su()["evaluation_dates"][-1])
    return unit, as_of, deck, "sec"


class TestEvalSnapshot(unittest.TestCase):
    def test_snapshot_roundtrip_is_identical(self):
        key = _key()
        fresh = S._compute_evaluation(*key)
        S._snapshot_put(*key, fresh)
        self.assertEqual(S._snapshot_get(*key), fresh)

    def test_stale_fingerprint_is_ignored(self):
        key = _key()
        S._snapshot_put(*key, S._compute_evaluation(*key))
        with mock.patch.object(S, "_eval_fingerprint", return_value="another-data-version"):
            self.assertIsNone(S._snapshot_get(*key))

    def test_restart_reads_snapshot_without_refitting(self):
        """模拟服务重启：进程内缓存清空后，评估直接读快照，不再逐井拟合。"""
        key = _key()
        expected = S._evaluate(*key)
        S.reset_cache()
        with mock.patch.object(S, "_compute_evaluation", side_effect=AssertionError("不应重算")):
            self.assertEqual(S._evaluate(*key), expected)
        S.reset_cache()


class TestUnitReport(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.unit = S.list_units()["plants"][0]["units"][0]["unit_id"]
        cls.tmp = tempfile.TemporaryDirectory()
        cls.out = R.build(cls.unit, formats=("docx", "html"), out_dir=Path(cls.tmp.name))

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_both_formats_written(self):
        for k in ("docx", "html"):
            p = Path(self.out[k])
            self.assertTrue(p.exists() and p.stat().st_size > 2000, k)
        self.assertTrue(self.out["trace_id"].startswith("rpt_"))

    def test_numbers_are_taken_from_services(self):
        from docx import Document
        comp = S.unit_sec_composition(self.unit, S._su()["evaluation_dates"][-1], "sec")
        doc = Document(self.out["docx"])
        text = "\n".join([p.text for p in doc.paragraphs]
                         + [c.text for t in doc.tables for row in t.rows for c in row.cells])
        html = Path(self.out["html"]).read_text(encoding="utf-8")
        for c in comp["components"] + [dict(reserves_t=comp["total_t"])]:
            self.assertIn(R.t(c["reserves_t"]), text)
            self.assertIn(R.t(c["reserves_t"]), html)
        self.assertIn("预评估", text)
        self.assertIn("附录 B", html)

    def test_unknown_scope_and_format_rejected(self):
        with self.assertRaises(S.KernelError):
            R.build("不存在的单元", formats=("html",), out_dir=Path(self.tmp.name))
        with self.assertRaises(S.KernelError):
            R.build(self.unit, formats=("pdf",), out_dir=Path(self.tmp.name))


if __name__ == "__main__":
    unittest.main()
