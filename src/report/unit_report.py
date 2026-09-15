"""SEC 单元储量预评估报告（Word 与可打印 HTML 两种格式）。

与界面、智能体同源：报告里的每个数字都取自 services 返回的 JSON，本模块只排版、不做结论性算术
（图件里的堆叠累计只是绘图坐标，不作为数字出现在正文）。
内容先组织成与格式无关的"块"，再分别渲染为 docx 与 html —— 两种格式的内容永远一致。
"""
from __future__ import annotations

import html
import io
import re
from datetime import datetime
from pathlib import Path
from typing import Callable, Dict, List, Optional, Sequence, Tuple

from .. import trace
from ..api import services as S
from ..config import path
from ..sec.checklist import DISCLAIMER

Block = Tuple[str, object]

COLORS = {"old_base": "#0E8F7F", "measure": "#C25E12", "new_infill": "#3B4FA8", "extension": "#B8487F"}
NEUTRAL, NEUTRAL_2, INK, INK_2, MUTED, RULE, GRID = ("#A9B1BC", "#6F7A89", "#1A2332", "#3F4A5A",
                                                    "#667185", "#DCE1E8", "#EBEEF2")
C1_SOFT = "#8CC7BE"
CAT_COLORS = {"PDP": "#1F7FB8", "PDNP": "#A67A0C", "PUD": "#7E57C2"}
CJK_FONTS = ("PingFang SC", "Hiragino Sans GB", "Noto Sans CJK SC", "Noto Sans SC", "Source Han Sans SC",
             "Microsoft YaHei", "SimHei", "WenQuanYi Micro Hei", "STHeiti", "Heiti SC", "Arial Unicode MS")
LEVEL_CN = {"unit": "SEC 单元", "plant": "采油厂", "company": "公司"}
RECON_SHORT = {"opening": "期初", "production": "产量消耗", "price_revision": "价格/成本", "new_wells": "提采新井",
               "measure": "措施", "extension": "扩边", "category_change": "类别调整",
               "technical_revision": "技术修订", "closing": "期末"}
RECON_COMP = {"new_wells": "new_infill", "measure": "measure", "extension": "extension"}
SERVICE_CN = {"comp": "SEC 储量构成", "prod": "逐月产量构成与计划对标", "decline": "老井基础递减",
              "new_wells": "新井识别", "measures": "措施效果", "reconcile": "储量对账",
              "attribution": "变化归因", "sensitivity": "敏感性分析", "indicators": "开发与经营指标",
              "categories": "证实储量类别", "tracking": "类别滚动", "depletion": "折耗与减值"}


# --------------------------------------------------------------------------- #
# 数值格式（只换单位与位数，不改变数值）
def t(v) -> str:
    if v is None:
        return "—"
    return f"{v / 1e4:,.2f} 万 t" if abs(v) >= 1e4 else f"{v:,.0f} t"


def st(v) -> str:
    return "—" if v is None else ("+" if v > 0 else "") + t(v)


def n(v, d: int = 0) -> str:
    return "—" if v is None else f"{v:,.{d}f}"


def wan(v) -> str:
    if v is None:
        return "—"
    return f"{v / 1e4:,.2f} 亿元" if abs(v) >= 1e4 else f"{v:,.0f} 万元"


def pct(v, d: int = 1) -> str:
    return "—" if v is None else f"{v:.{d}f}%"


def _compact(v) -> str:
    if v is None:
        return ""
    return f"{v / 1e4:.1f}万" if abs(v) >= 1e4 else f"{v:,.0f}"


def _evidence(kind: str, e: Dict) -> str:
    if kind == "wells":
        return f"{e['well_code']}（{e['unit_id']}）{t(e.get('reserves_t'))}，{e.get('basis')}，投产 {e.get('months_on')} 个月"
    if kind == "measures":
        return f"{e['well_code']}（{e['unit_id']}）{e.get('event_name')} {e.get('event_ym')}，增加可采 {t(e.get('inc_eur_t'))}"
    if kind == "economics":
        return (f"油价 {n(e['price_open_usd_bbl'], 2)} → {n(e['price_close_usd_bbl'], 2)} USD/bbl，"
                f"单井经济极限 {n(e['q_econ_open_t_per_d'], 2)} → {n(e['q_econ_close_t_per_d'], 2)} t/d")
    if kind == "wells_rate_change":
        return (f"{e['well_code']}（{e['unit_id']}）日产 {n(e['rate_open_t_per_d'], 2)} → "
                f"{n(e['rate_close_t_per_d'], 2)} t/d，含水 {pct(e.get('water_cut_open_pct'))} → "
                f"{pct(e.get('water_cut_close_pct'))}")
    return ""


# --------------------------------------------------------------------------- #
# 取数
def gather(scope: str, as_of: Optional[str] = None, scenario: str = "sec") -> Dict:
    units = S.list_units()
    dates = units["evaluation_dates"]
    as_of = str(as_of or dates[-1])
    comp = S.unit_sec_composition(scope, as_of, scenario)          # 评估对象 / 基准日不合法在这里直接报错
    prev = dates[dates.index(as_of) - 1] if dates.index(as_of) > 0 else None
    data: Dict = dict(units=units, scope=scope, as_of=as_of, scenario=scenario, prev=prev, comp=comp)
    no_prev = {"_error": f"{as_of} 是首个评估基准日，没有上期结果可对账"}
    jobs: List[Tuple[str, Callable, tuple, dict]] = [
        ("prod", S.unit_production_composition, (scope, as_of), {}),
        ("decline", S.unit_base_decline, (scope, as_of), {}),
        ("new_wells", S.unit_new_wells, (scope, as_of), {}),
        ("measures", S.unit_measure_effects, (scope, as_of), {}),
        ("sensitivity", S.unit_sensitivity, (scope, as_of, scenario), {}),
        ("indicators", S.unit_indicators, (scope,), dict(as_of_list=",".join(d for d in (prev, as_of) if d))),
        ("categories", S.unit_proved_categories, (scope, as_of, scenario), {}),
        ("depletion", S.unit_depletion_impairment, (scope, as_of), {}),
    ]
    if prev:
        jobs += [("reconcile", S.unit_reconcile, (scope,), dict(from_as_of=prev, to_as_of=as_of, scenario=scenario)),
                 ("attribution", S.unit_change_attribution, (scope,),
                  dict(from_as_of=prev, to_as_of=as_of, scenario=scenario)),
                 ("tracking", S.unit_category_tracking, (scope,), dict(from_as_of=prev, to_as_of=as_of, scenario=scenario)),
                 ("pud_disclosure", S.unit_pud_disclosure, (scope,), dict(from_as_of=prev, to_as_of=as_of, scenario=scenario))]
    else:
        data["reconcile"] = data["attribution"] = data["tracking"] = data["pud_disclosure"] = no_prev
    for key, fn, args, kw in jobs:
        try:
            data[key] = fn(*args, **kw)
        except S.KernelError as exc:
            data[key] = {"_error": str(exc)}          # 失败如实写进报告，不留白也不编造
    data["scenarios"] = []
    for sc in units["scenarios"]:
        try:
            data["scenarios"].append(S.unit_sec_composition(scope, as_of, sc["key"]))
        except S.KernelError as exc:
            data["scenarios"].append({"_error": str(exc), "scenario_label": sc["label"]})
    return data


# --------------------------------------------------------------------------- #
# 图件（matplotlib，无中文字体时如实省略）
def _mpl():
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        from matplotlib import font_manager
    except ImportError:
        return None
    names = {f.name for f in font_manager.fontManager.ttflist}
    cjk = [f for f in CJK_FONTS if f in names]
    if not cjk:
        return None
    plt.rcParams.update({
        "font.family": "sans-serif", "font.sans-serif": cjk + ["DejaVu Sans"], "axes.unicode_minus": False,
        "svg.fonttype": "none", "font.size": 8.5, "axes.edgecolor": RULE, "axes.labelcolor": MUTED,
        "xtick.color": MUTED, "ytick.color": MUTED, "axes.spines.top": False, "axes.spines.right": False,
        "axes.grid": True, "grid.color": GRID, "grid.linewidth": 0.8, "axes.axisbelow": True,
        "legend.frameon": False, "xtick.major.size": 0, "ytick.major.size": 0})
    return plt


def _wan(ax, axis: str = "y") -> None:
    from matplotlib.ticker import FuncFormatter
    f = FuncFormatter(lambda v, _: f"{v / 1e4:g}万" if abs(v) >= 1e4 else f"{v:,.0f}")
    (ax.yaxis if axis == "y" else ax.xaxis).set_major_formatter(f)


def _export(fig, plt) -> Dict[str, object]:
    png, svg = io.BytesIO(), io.BytesIO()
    fig.savefig(png, format="png", dpi=200, bbox_inches="tight", facecolor="white")
    fig.savefig(svg, format="svg", bbox_inches="tight", facecolor="white")
    plt.close(fig)
    s = svg.getvalue().decode("utf-8")
    return dict(png=png.getvalue(), svg=s[s.index("<svg"):])


def _fig_units(plt, comp: Dict):
    units = comp["units"]
    names = {c["key"]: c["name"] for c in comp["components"]}
    fig, ax = plt.subplots(figsize=(7.2, 0.34 * len(units) + 1.0))
    y = list(range(len(units)))
    left = [0.0] * len(units)
    for k, color in COLORS.items():
        vals = [max(u["components"].get(k) or 0.0, 0.0) for u in units]
        ax.barh(y, vals, left=left, height=0.56, color=color, edgecolor="white", linewidth=1.2, label=names.get(k, k))
        left = [a + b for a, b in zip(left, vals)]
    for yi, u, x in zip(y, units, left):
        ax.text(x, yi, "  " + _compact(u["total_t"]), va="center", fontsize=8, color=MUTED)
    ax.set_yticks(y, [u["unit_id"] for u in units])
    ax.invert_yaxis()
    ax.grid(axis="y", visible=False)
    ax.margins(x=0.1)
    _wan(ax, "x")
    ax.legend(ncol=4, loc="lower left", bbox_to_anchor=(0, 1.0), fontsize=8, handlelength=1, columnspacing=1.2)
    return fig


def _fig_categories(plt, cat: Dict):
    units = cat["units"]
    fig, ax = plt.subplots(figsize=(7.2, 0.34 * len(units) + 1.0))
    y = list(range(len(units)))
    left = [0.0] * len(units)
    names = {c["key"]: c["name"] for c in cat["categories"]}
    for k, color in CAT_COLORS.items():
        vals = [max(u.get(k) or 0.0, 0.0) for u in units]
        ax.barh(y, vals, left=left, height=0.56, color=color, edgecolor="white", linewidth=1.2, label=names.get(k, k))
        left = [a + b for a, b in zip(left, vals)]
    for yi, u, x in zip(y, units, left):
        ax.text(x, yi, "  " + _compact(u["total_t"]), va="center", fontsize=8, color=MUTED)
    ax.set_yticks(y, [u["unit_id"] for u in units])
    ax.invert_yaxis()
    ax.grid(axis="y", visible=False)
    ax.margins(x=0.1)
    _wan(ax, "x")
    ax.legend(ncol=3, loc="lower left", bbox_to_anchor=(0, 1.0), fontsize=8, handlelength=1)
    return fig


ROLL_TINY = {"opening": "期初", "converted": "钻井转开发", "expired_5yr": "五年规则移出", "removed": "其他移出",
             "revision": "修订", "new_booking": "新入账", "closing": "期末"}


def _fig_roll(plt, rows: List[Dict]):
    fig, ax = plt.subplots(figsize=(7.2, 2.8))
    run, bars = 0.0, []
    for r in rows:
        total = r["key"] in ("opening", "closing")
        a, b = (0.0, r["value"]) if total else (run, run + r["value"])
        run = b
        bars.append((r, a, b, total))
    for i, (r, a, b, total) in enumerate(bars):
        color = NEUTRAL if total else CAT_COLORS["PDP"] if r["key"] == "converted" else \
            CAT_COLORS["PUD"] if r["key"] == "new_booking" else NEUTRAL_2
        ax.bar(i, b - a, bottom=a, width=0.58, color=color)
        if i < len(bars) - 1:
            ax.plot([i + 0.29, i + 0.71], [b, b], color=RULE, lw=0.8)
        ax.text(i, max(a, b), ("" if total or r["value"] < 0 else "+") + _compact(r["value"]), ha="center", va="bottom",
                fontsize=7.5, color=MUTED)
    ax.set_xticks(range(len(bars)), [ROLL_TINY.get(r["key"], r["item"]) for r, *_ in bars], fontsize=7.5)
    ax.axhline(0, color=RULE, lw=0.8)
    ax.grid(axis="x", visible=False)
    _wan(ax)
    return fig


def _fig_monthly(plt, prod: Dict):
    months = prod["months"]
    x = list(range(len(months)))
    fig, ax = plt.subplots(figsize=(7.2, 2.9))
    bottom = [0.0] * len(months)
    for key, label, color in (("old_base_t", "老井基础", COLORS["old_base"]),
                              ("measure_inc_t", "本期措施增油", COLORS["measure"]),
                              ("new_oil_t", "本期新井（提采 + 扩边）", COLORS["new_infill"])):
        vals = [m[key] or 0.0 for m in months]
        ax.bar(x, vals, bottom=bottom, width=0.78, color=color, edgecolor="white", linewidth=0.5, label=label)
        bottom = [a + b for a, b in zip(bottom, vals)]
    ax.set_xticks(x[::6], [m["ym"] for m in months][::6])
    ax.grid(axis="x", visible=False)
    ax.set_ylabel("月产油 (t)")
    _wan(ax)
    ax.legend(ncol=3, loc="lower left", bbox_to_anchor=(0, 1.0), fontsize=8, handlelength=1)
    return fig


def _fig_waterfall(plt, recon: Dict):
    rows = recon["table"]
    fig, ax = plt.subplots(figsize=(7.2, 3.0))
    run, bars = 0.0, []
    for r in rows:
        total = r["key"] in ("opening", "closing")
        a, b = (0.0, r["value"]) if total else (run, run + r["value"])
        run = b
        bars.append((r, a, b, total))
    big = sorted((x for x in bars if not x[3]), key=lambda x: -abs(x[0]["value"]))[:2]
    for i, (r, a, b, total) in enumerate(bars):
        color = NEUTRAL if total else COLORS.get(RECON_COMP.get(r["key"], ""), NEUTRAL_2)
        ax.bar(i, b - a, bottom=a, width=0.58, color=color)
        if i < len(bars) - 1:
            ax.plot([i + 0.29, i + 0.71], [b, b], color=RULE, lw=0.8)
        if total or (r, a, b, total) in big:
            ax.text(i, max(a, b), ("" if total or r["value"] < 0 else "+") + _compact(r["value"]),
                    ha="center", va="bottom", fontsize=7.5, color=MUTED)
    ax.set_xticks(range(len(bars)), [RECON_SHORT.get(r["key"], r["item"]) for r, *_ in bars], fontsize=7.5)
    ax.axhline(0, color=RULE, lw=0.8)
    ax.grid(axis="x", visible=False)
    _wan(ax)
    return fig


def _fig_sensitivity(plt, sens: Dict):
    curves = sens["curves"]
    fig, axes = plt.subplots(2, 2, figsize=(7.2, 4.4))
    base = {"price": sens.get("base_price_usd_bbl"), "opex": 1.0, "rate": 0.0, "decline": 0.0}
    for ax, cv in zip(axes.flat, curves):
        ax.axhline(0, color=RULE, lw=0.8)
        if base.get(cv["param"]) is not None:
            ax.axvline(base[cv["param"]], color=RULE, lw=0.8, ls="--")
        ax.plot(cv["x"], cv["delta_t"], color=INK_2, lw=1.8, marker="o", ms=2.6)
        ax.set_title(cv["name"], fontsize=9, loc="left", color=INK)
        ax.set_xlabel(cv["unit"], fontsize=8)
        _wan(ax)
    for ax in list(axes.flat)[len(curves):]:
        ax.axis("off")
    fig.tight_layout()
    return fig


def _fig_indicators(plt, indi: Dict):
    from matplotlib.lines import Line2D
    pa, pb = indi["periods"][0], indi["periods"][-1]
    rows = pb["indicators"]
    prev = {r["key"]: r for r in pa["indicators"]}
    fig, ax = plt.subplots(figsize=(7.2, 0.28 * len(rows) + 0.9))
    for i, r in enumerate(rows):
        a, b = prev.get(r["key"], {}).get("score"), r["score"]
        if a is not None and b is not None:
            ax.plot([a, b], [i, i], color=RULE, lw=2, zorder=1)
        if a is not None:
            ax.scatter([a], [i], s=34, color=C1_SOFT, zorder=2, edgecolors="white", linewidths=1)
        if b is not None:
            ax.scatter([b], [i], s=34, color=COLORS["old_base"], zorder=3, edgecolors="white", linewidths=1)
    ax.set_yticks(range(len(rows)), [r["name"] for r in rows])
    ax.invert_yaxis()
    ax.set_xlim(-3, 103)
    ax.set_xlabel("得分（0 ~ 100）")
    ax.grid(axis="y", visible=False)
    ax.legend(handles=[Line2D([], [], marker="o", ls="", color=C1_SOFT, label=pa["as_of"]),
                       Line2D([], [], marker="o", ls="", color=COLORS["old_base"], label=pb["as_of"])],
              ncol=2, loc="lower left", bbox_to_anchor=(0, 1.0), fontsize=8)
    return fig


# --------------------------------------------------------------------------- #
# 内容块
def _table(cols: Sequence[str], rows: List[List[str]], num: Sequence[bool], total: bool = False,
           wrap: Sequence[int] = ()) -> Block:
    """wrap：允许折行的列（长文字说明）；其余列不折行，井号、年月不会被拆成两行。"""
    return ("table", dict(cols=list(cols), rows=rows, num=list(num), total=total, wrap=list(wrap)))


def build_blocks(data: Dict, report_trace_id: str) -> List[Block]:
    comp, as_of, prev = data["comp"], data["as_of"], data["prev"]
    sc = comp["scope"]
    plt = _mpl()
    B: List[Block] = []

    def figure(name: str, caption: str, make: Callable) -> None:
        if plt is None:
            B.append(("note", f"图件“{caption}”未生成：运行环境缺少 matplotlib 或中文字体。"))
            return
        try:
            B.append(("figure", dict(caption=caption, **_export(make(plt), plt))))
        except Exception as exc:                     # 图件失败不拖垮整份报告，如实注明
            B.append(("note", f"图件“{caption}”生成失败：{exc}"))

    def err(key: str) -> Optional[str]:
        v = data.get(key) or {}
        return v.get("_error") if isinstance(v, dict) else None

    # 封面
    B.append(("title", dict(title=f"{sc['name']} SEC 储量预评估报告",
                            subtitle=f"基准日 {as_of} · {comp['scenario_label']}口径 · 初稿")))
    B.append(("meta", [
        ("评估对象", f"{sc['name']}（{LEVEL_CN.get(sc['level'], sc['level'])}，{sc['n_units']} 个 SEC 单元）"),
        ("评估基准日", f"{as_of}（本期 {comp['period_start_ym']} ~ {comp['period_end_ym']}）"),
        ("价格情景", f"{comp['scenario_label']}｜价格册 {comp['price_deck_id']}｜油价 {n(comp['economics']['price_usd_bbl'], 2)} USD/bbl"),
        ("单井经济极限", f"{n(comp['economics']['q_econ_t_per_d'], 2)} t/d（{comp['economics']['note']}）"),
        ("数据来源", comp["data_source"]),
        ("模型版本 / 标签口径", f"{comp['model_version']} / {comp['label_def_version']}"),
        ("生成时间", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("报告追溯号", report_trace_id),
    ]))
    B.append(("disclaimer", DISCLAIMER))
    B.append(("pagebreak", None))

    # 一、摘要
    B.append(("h1", "一、评估结论摘要"))
    parts = "、".join(f"{c['name']} {t(c['reserves_t'])}（{pct(c['share_pct'])}）" for c in comp["components"])
    bullets = [f"{sc['name']}在 {as_of}、{comp['scenario_label']}口径下，已证实已开发储量（PDP）为 {t(comp['total_t'])}，"
               f"其中{parts}。"]
    nwc, ms = comp["new_wells"], comp["measures"]
    bullets.append(f"本期（{comp['period_start_ym']} ~ {comp['period_end_ym']}）产油 {t(comp['production_in_period_t'])}；"
                   f"投产提采新井 {nwc['n_infill']} 口、扩边井 {nwc['n_extension']} 口；实施措施 {ms['n_in_period']} 次，"
                   f"有效 {ms['n_effective']} 次，增加可采储量 {t(ms['inc_eur_total_t'])}。")
    if not err("attribution") and not err("reconcile"):
        att, rc = data["attribution"], data["reconcile"]
        top = "、".join(f"{d['item']} {st(d['value_t'])}（{pct(d['share_pct'])}）" for d in att["drivers"][:3])
        bullets.append(f"与 {att['from_as_of']} 相比，PDP 由 {t(att['opening_t'])} 变为 {t(att['closing_t'])}，"
                       f"变化 {st(att['change_t'])}；对账{'闭合' if rc['balanced'] else '未闭合（差额 ' + n(rc['difference'], 1) + ' t）'}，"
                       f"产量法折耗率 {pct(rc['depletion_rate_pct'], 2)}。影响最大的三项：{top}。")
    if not err("categories"):
        cat = data["categories"]
        by = {c["key"]: c for c in cat["categories"]}
        conv = ""
        if not err("tracking") and data["tracking"].get("pud"):
            conv = f"；本期 PUD 转化率 {pct(data['tracking']['pud']['conversion_rate_pct'])}"
        bullets.append(f"证实储量合计 {t(cat['total_proved_t'])}：PDP {t(by['PDP']['reserves_t'])}、"
                       f"PDNP {t(by['PDNP']['reserves_t'])}（{by['PDNP']['n_items']} 口停产井）、"
                       f"PUD {t(by['PUD']['reserves_t'])}（{by['PUD']['n_items']} 个部署井位）{conv}。")
    if not err("depletion"):
        sm = data["depletion"]["summary"]
        bullets.append(f"产量法折耗率 {pct(sm['depletion_rate_pct'], 2)}，折耗额 {wan(sm['depletion_wan'])}；"
                       f"减值测试 {sm['n_impaired']} 个单元减值，减值额 {wan(sm['impairment_wan'])}，期末资产净值 {wan(sm['closing_nbv_wan'])}。")
    scs = [s for s in data["scenarios"] if "_error" not in s]
    if scs:
        bullets.append("三种价格情景下的 PDP：" + "；".join(f"{s['scenario_label']} {t(s['total_t'])}" for s in scs) + "。")
    if not err("sensitivity"):
        w = max(data["sensitivity"]["weights"], key=lambda x: x["weight_pct"] or 0)
        bullets.append(f"对 PDP 最敏感的参数是{w['name']}（敏感权重 {pct(w['weight_pct'])}，{data['sensitivity']['perturbation']}）。")
    if not err("indicators"):
        ps = data["indicators"]["periods"]
        groups = []
        for g in ps[-1]["groups"]:
            if len(ps) > 1:
                g0 = next((x for x in ps[0]["groups"] if x["group"] == g["group"]), {})
                groups.append(f"{g['name']}综合分 {n(g0.get('score'), 1)}（{ps[0]['as_of']}）→ {n(g['score'], 1)}（{ps[-1]['as_of']}）")
            else:
                groups.append(f"{g['name']}综合分 {n(g['score'], 1)}")
        bullets.append("；".join(groups) + "。")
    B.append(("bullets", bullets))

    # 二、PDP 构成
    B.append(("h1", "二、PDP 储量构成"))
    B.append(_table(["构成", "剩余可采", "占比", "数量", "取值依据"],
                    [[c["name"], t(c["reserves_t"]), pct(c["share_pct"]),
                      f"{c['n_items']} {'次' if c['key'] == 'measure' else '口'}", c["basis"]] for c in comp["components"]]
                    + [["合计", t(comp["total_t"]), "100.0%", "", ""]],
                    [False, True, True, True, False], total=True, wrap=[4]))
    if len(comp["units"]) > 1:
        figure("units", "各单元 PDP 构成", lambda p: _fig_units(p, comp))
        B.append(_table(["单元", "名称"] + [c["name"] for c in comp["components"]] + ["合计"],
                        [[u["unit_id"], u["unit_name"]] + [t(u["components"].get(c["key"])) for c in comp["components"]]
                         + [t(u["total_t"])] for u in comp["units"]],
                        [False, False] + [True] * (len(comp["components"]) + 1)))
    ow = comp["old_wells"]
    B.append(("note", f"老井基础：参与评估 {ow['n_evaluated']} 口（其中峰后历史不足、按类比法取值 {ow['n_short_history']} 口，"
                      f"以措施前基线外推 {ow['n_measure_baseline']} 口），近期无产量不计入 PDP {ow['n_not_producing']} 口，"
                      f"资料不足 {ow['n_insufficient']} 口；最佳估计 {t(ow['best_estimate_t'])}。"
                      + (comp["aggregation_note"] or "")))

    # 三、价格情景
    B.append(("h1", "三、价格情景对比"))
    B.append(_table(["情景", "油价 (USD/bbl)", "单井经济极限 (t/d)", "PDP"] + [c["name"] for c in comp["components"]],
                    [[s["scenario_label"], "—", "—", s["_error"]] + ["—"] * len(comp["components"]) if "_error" in s else
                     [s["scenario_label"], n(s["economics"]["price_usd_bbl"], 2), n(s["economics"]["q_econ_t_per_d"], 2),
                      t(s["total_t"])] + [t(c["reserves_t"]) for c in s["components"]] for s in data["scenarios"]],
                    [False, True, True, True] + [True] * len(comp["components"])))

    # 四、产量与递减
    B.append(("h1", "四、产量构成与老井递减"))
    if err("prod"):
        B.append(("p", "逐月产量构成不可得：" + err("prod")))
    else:
        prod = data["prod"]
        figure("monthly", f"逐月“新-老-措”产量构成（本期 {prod['period_start_ym']} ~ {prod['period_end_ym']}）",
               lambda p: _fig_monthly(p, prod))
        B.append(("h2", "（一）计划与实际对标"))
        PT = [("new_wells", "新井投产（口）"), ("new_oil", "新井产油（t）"), ("measure_wells", "措施（井次）"),
              ("measure_inc", "措施增油（t）"), ("old_oil", "老井产油（t）")]
        B.append(_table(["项目", "计划", "实际", "偏差"],
                        [[label, n(prod["period_totals"][k]["plan"]), n(prod["period_totals"][k]["actual"]),
                          pct(prod["period_totals"][k]["deviation_pct"])] for k, label in PT],
                        [False, True, True, True]))
        if not prod["plan_available"]:
            B.append(("note", "本期未导入计划数据。"))
        B.append(("note", prod["note"] + "。"))
    if err("decline"):
        B.append(("p", "老井基础递减不可得：" + err("decline")))
    else:
        dec = data["decline"]
        B.append(("h2", "（二）老井基础递减率"))
        B.append(_table(["扣除年限", "统计窗口", "自然递减率（年）", "自然递减率（月）", "拟合 R²", "综合递减率（年）", "综合递减率（月）"],
                        [[f"扣近 {r['exclude_years']} 年", f"{r['window_start_ym']} ~ {r['window_end_ym']}",
                          pct(r["natural_annual_pct"], 2), pct(r["natural_monthly_pct"], 3), n(r["natural_fit_r2"], 3),
                          pct(r["comprehensive_annual_pct"], 2), pct(r["comprehensive_monthly_pct"], 3)] for r in dec["results"]],
                        [False, False, True, True, True, True, True]))
        B.append(("note", dec["definition"] + "。"))

    # 五、新井
    B.append(("h1", "五、本期新井识别"))
    if err("new_wells"):
        B.append(("p", "新井识别不可得：" + err("new_wells")))
    else:
        nw = data["new_wells"]
        basis = "、".join(f"{k} {v} 口" for k, v in (nwc.get("basis_counts") or {}).items())
        B.append(("p", f"识别规则：{nw['rule']['text']}。本期识别提采新井 {nw['n_infill']} 口、扩边井 {nw['n_extension']} 口"
                       + (f"；储量取值方法：{basis}" if basis else "") + "。完整清单见附录 A。"))
        for code, label in (("infill", "提采新井"), ("extension", "扩边井")):
            rows = [w for w in nw["wells"] if w["category_code"] == code][:8]
            if rows:
                B.append(("h2", f"{label}（剩余可采前 {len(rows)} 口）"))
                B.append(_table(["井号", "单元", "投产月", "周边老井", "最近老井 (m)", "剩余可采", "取值方法"],
                                [[w["well_code"], w["unit_id"], w["first_prod_ym"], n(w["n_old_neighbors"]),
                                  n(w["nearest_old_m"]), t(w["reserves_t"]), w["basis"]] for w in rows],
                                [False, False, False, True, True, True, False]))

    # 六、措施
    B.append(("h1", "六、本期措施效果"))
    if err("measures"):
        B.append(("p", "措施效果不可得：" + err("measures")))
    else:
        me = data["measures"]
        o = me["overall"]
        B.append(("p", f"本期措施 {o['n']} 次，有效 {o['n_effective']} 次（有效率 {pct(o['effective_rate_pct'])}）；"
                       f"单井日产由措施前 {n(o['pre_rate_avg_t_per_d'], 2)} t/d 提高到 {n(o['post_rate_avg_t_per_d'], 2)} t/d；"
                       f"增加可采储量合计 {t(o['inc_eur_total_t'])}。"))
        B.append(_table(["措施类型", "次数", "有效", "措施前日产 (t/d)", "措施后日产 (t/d)", "单次增加可采", "增加可采合计"],
                        [[r["name"], n(r["n"]), n(r["n_effective"]), n(r["pre_rate_avg_t_per_d"], 2),
                          n(r["post_rate_avg_t_per_d"], 2), t(r["inc_eur_avg_t"]), t(r["inc_eur_total_t"])]
                         for r in sorted(me["by_type"], key=lambda r: -(r["inc_eur_total_t"] or 0))],
                        [False, True, True, True, True, True, True]))
        B.append(("note", me["definition"] + "。"))

    # 七、对账
    B.append(("h1", "七、储量对账与变化归因"))
    if err("reconcile"):
        B.append(("p", "储量对账不可得：" + err("reconcile")))
    else:
        rc = data["reconcile"]
        B.append(("p", f"{rc['from_as_of']} → {rc['to_as_of']}，期初来源：{rc['opening_source']}；"
                       f"对账{'闭合' if rc['balanced'] else '未闭合'}（差额 {n(rc['difference'], 1)} t）；"
                       f"产量法折耗率 {pct(rc['depletion_rate_pct'], 2)}。"))
        figure("waterfall", "储量对账瀑布图", lambda p: _fig_waterfall(p, rc))
        B.append(_table(["行项", "储量变动", "说明"],
                        [[r["item"], st(r["value"]) if r["key"] not in ("opening", "closing") else t(r["value"]), r["note"] or ""]
                         for r in rc["table"]], [False, True, False], wrap=[2]))
        B.append(("note", f"{rc['note']}。{rc['depletion_note']}。"))
    if not err("attribution"):
        att = data["attribution"]
        B.append(("h2", "变化归因（按影响大小排序）"))
        B.append(("bullets", [f"{d['item']} {st(d['value_t'])}（占 {pct(d['share_pct'])}）："
                              + ("；".join(_evidence(d["evidence_kind"], e) for e in d["evidence"]) or "—")
                              for d in att["drivers"]]))
        B.append(("note", att["note"] + "。"))

    # 八、证实储量类别
    B.append(("h1", "八、证实储量类别（PDP / PDNP / PUD）"))
    if err("categories"):
        B.append(("p", "证实储量类别不可得：" + err("categories")))
    else:
        cat = data["categories"]
        B.append(("p", f"{cat['scenario_label']}口径下证实储量合计 {t(cat['total_proved_t'])}。"))
        if len(cat["units"]) > 1:
            figure("categories", "各单元证实储量类别", lambda p: _fig_categories(p, cat))
        B.append(_table(["类别", "储量", "占比", "数量", "取值依据"],
                        [[c["name"], t(c["reserves_t"]), pct(c["share_pct"]),
                          f"{c['n_items']} {'个井位' if c['key'] == 'PUD' else '口井'}", c["basis"]] for c in cat["categories"]]
                        + [["合计", t(cat["total_proved_t"]), "100.0%", "", ""]],
                        [False, True, True, True, False], total=True, wrap=[4]))
        pu = cat["pud"]
        B.append(("h2", "（一）PUD 部署井位"))
        B.append(("p", f"入账规则：{pu['rule']}。本期井位判定：" + "，".join(
            f"{S.PUD_STATUS_CN.get(k, k)} {v} 个" for k, v in pu["n_by_status"].items()) + "。"))
        booked = [l for l in pu["locations"] if l["status"] == "booked"]
        if booked:
            B.append(_table(["井位", "单元", "计划钻井", "首次入账", "五年期限", "周边在产井", "剩余可采"],
                            [[l["location_id"], l["unit_id"], l["planned_drill_ym"], l["first_booked_as_of"] or "本期",
                              l["deadline_ym"], n(l["n_producing_neighbors"]), t(l["reserves_t"])] for l in booked[:15]],
                            [False, False, False, False, False, True, True]))
        if cat["warnings"]:
            B.append(("bullets", ["预警：" + w["text"] for w in cat["warnings"]]))
        tr = None if err("tracking") else data["tracking"]
        if tr and tr.get("pud"):
            tp = tr["pud"]
            B.append(("p", f"{tr['from_as_of']} → {tr['to_as_of']}：期初入账 {tp['n_opening']} 个井位，钻井转化 {tp['n_converted']} 个"
                           f"（转化率 {pct(tp['conversion_rate_pct'])}），五年规则移出 {tp['n_expired']} 个，其他移出 {tp['n_removed']} 个，"
                           f"新入账 {tp['n_new']} 个，期末 {tp['n_closing']} 个。"))
            figure("pud_roll", "PUD 滚动", lambda p: _fig_roll(p, tp["table"]))
            B.append(_table(["行项", "储量变动"], [[r["item"], t(r["value"]) if r["key"] in ("opening", "closing") else st(r["value"])]
                                               for r in tp["table"]], [False, True]))
        dsc = None if err("pud_disclosure") else data["pud_disclosure"]
        if dsc:
            B.append(("h2", "（附）Item 1203 已证实未开发储量披露草稿"))
            for sec in dsc["sections"]:
                B.append(("p", f"{sec['item']} {sec['title']}：{sec['text']}（依据：{'；'.join(sec['citations'])}）"))
            B.append(("note", dsc["draft_note"]))
        pdn = cat["pdnp"]
        B.append(("h2", "（二）PDNP 停产井"))
        B.append(("p", f"规则：{pdn['rule']}。停产井判定：" + "，".join(
            f"{S.PDNP_STATUS_CN.get(k, k)} {v} 口" for k, v in pdn["n_by_status"].items())
                       + f"；计入 PDNP {pdn['n_booked']} 口，{t(pdn['reserves_t'])}。"))
        bw = [w for w in pdn["wells"] if w["status"] == "booked"][:15]
        if bw:
            B.append(_table(["井号", "单元", "最后生产", "停产月数", "PDNP 储量", "停产时日产 (t/d)"],
                            [[w["well_code"], w["unit_id"], w["last_prod_ym"], n(w["shut_in_months"]), t(w["reserves_t"]),
                              n(w["rate_at_shut_t_per_d"], 2)] for w in bw], [False, False, False, True, True, True]))
        if tr and tr.get("pdnp"):
            td, cc = tr["pdnp"], tr["pdp_category_change"]
            B.append(("p", f"期初 PDNP {td['n_opening']} 口，复产转 PDP {td['n_reactivated']} 口，长停或不经济移出 {td['n_removed']} 口，"
                           f"新增停产 {td['n_new']} 口，期末 {td['n_closing']} 口。PDP 对账中的类别调整：停产井复产转入 "
                           f"{t(cc['reactivated_t'])}（{cc['n_reactivated']} 口），在产井停井转出 {st(cc['shut_in_t'])}（{cc['n_shut_in']} 口）。"))
        B.append(("note", (tr or cat)["note"] + "。"))

    # 九、折耗与减值
    B.append(("h1", "九、折耗与减值测试"))
    if err("depletion"):
        B.append(("p", "折耗与减值不可得：" + err("depletion")))
    else:
        dp = data["depletion"]
        sm, asm = dp["summary"], dp["assumptions"]
        B.append(("p", f"期初资产净值 {wan(sm['opening_nbv_wan'])}，本期资本化投入 {wan(sm['capex_additions_wan'])}；"
                       f"产量法折耗率 {pct(sm['depletion_rate_pct'], 2)}，折耗额 {wan(sm['depletion_wan'])}，折耗后账面价值 {wan(sm['carrying_wan'])}。"
                       f"按{asm['impairment_scenario']}、折现率 {asm['discount_rate']:.0%} 测算可收回金额 {wan(sm['recoverable_wan'])}，"
                       f"{sm['n_impaired']} 个单元发生减值，减值额 {wan(sm['impairment_wan'])}，期末资产净值 {wan(sm['closing_nbv_wan'])}。"))
        B.append(_table(["单元", "期初净值", "本期投入", "折耗率", "折耗额", "账面价值", "可收回金额", "减值额", "期末净值"],
                        [[u["unit_id"], wan(u["opening_nbv_wan"]), wan(u["capex_additions_wan"]), pct(u["depletion_rate_pct"], 2),
                          wan(u["depletion_wan"]), wan(u["carrying_wan"]), wan(u["recoverable_wan"]),
                          wan(u["impairment_wan"]) if u["impaired"] else "—", wan(u["closing_nbv_wan"])] for u in dp["units"]],
                        [False] + [True] * 8))
        B.append(("bullets", dp["method"]))
        B.append(("note", dp["note"] + f"。汇率 {asm['fx_cny_per_usd']} 元/美元。"))

    # 十、敏感性
    B.append(("h1", "十、敏感性分析"))
    if err("sensitivity"):
        B.append(("p", "敏感性分析不可得：" + err("sensitivity")))
    else:
        se = data["sensitivity"]
        B.append(("p", f"{se['scenario_label']}，基准油价 {n(se['base_price_usd_bbl'], 2)} USD/bbl，"
                       f"最佳估计 {t(se['base_best_estimate_t'])}；扰动幅度：{se['perturbation']}。"))
        figure("sensitivity", "各参数对 PDP 的影响（相对基准的变化量）", lambda p: _fig_sensitivity(p, se))
        B.append(_table(["参数", "PDP 摆幅", "敏感权重"],
                        [[w["name"], t(w["swing_t"]), pct(w["weight_pct"])] for w in se["weights"]], [False, True, True]))
        if len(se["units"]) > 1:
            names = [w["name"] for w in se["weights"]]
            B.append(("h2", "各单元参数敏感权重"))
            B.append(_table(["单元", "最敏感参数"] + names,
                            [[u["unit_id"], u["top_param"]] + [pct(w["weight_pct"]) for w in u["weights"]] for u in se["units"]],
                            [False, False] + [True] * len(names)))
        B.append(("note", se["note"] + "。"))

    # 十一、指标
    B.append(("h1", "十一、开发与经营指标评价"))
    if err("indicators"):
        B.append(("p", "指标评价不可得：" + err("indicators")))
    else:
        ind = data["indicators"]
        ps = ind["periods"]
        if len(ps) > 1:
            figure("indicators", f"开发与经营指标得分两期对比（{ps[0]['as_of']} / {ps[-1]['as_of']}）",
                   lambda p: _fig_indicators(p, ind))
        cols = ["指标", "单位"]
        for p in ps:
            cols += [f"{p['as_of']} 值", "得分"]
        rows = []
        for r in ps[-1]["indicators"]:
            row = [r["name"], r["unit"]]
            for p in ps:
                x = next((y for y in p["indicators"] if y["key"] == r["key"]), {})
                row += [n(x.get("value"), 2), n(x.get("score"), 1)]
            rows.append(row)
        B.append(_table(cols, rows, [False, False] + [True] * (2 * len(ps))))
        B.append(("note", ind["scoring"] + "。"))

    # 十二、口径与局限
    B.append(("h1", "十二、口径与局限"))
    lim = [DISCLAIMER, "报告中的全部数值均取自平台算法内核的返回结果，与系统界面、智能问答同源，可凭附录 B 的追溯号回查。"]
    if comp.get("aggregation_note"):
        lim.append(comp["aggregation_note"] + "。")
    if comp["data_source"] != "REAL":
        lim.append(f"本报告基于 {comp['data_source']} 数据生成，仅用于方法演示，不代表任何真实油田的储量。")
    B.append(("bullets", lim))

    # 附录
    if not err("new_wells") and data["new_wells"]["wells"]:
        B.append(("pagebreak", None))
        B.append(("h1", "附录 A　本期新井清单"))
        B.append(_table(["井号", "单元", "分类", "投产月", "周边老井", "最近老井 (m)", "剩余可采", "取值方法"],
                        [[w["well_code"], w["unit_id"], w["category"], w["first_prod_ym"],
                          n(w["n_old_neighbors"]), n(w["nearest_old_m"]), t(w["reserves_t"]), w["basis"]]
                         for w in data["new_wells"]["wells"]],
                        [False, False, False, False, True, True, True, False]))
    B.append(("h1", "附录 B　数据追溯"))
    trace_rows = [["报告", report_trace_id]]
    for key, label in SERVICE_CN.items():
        v = data.get(key) or {}
        trace_rows.append([label, v.get("trace_id") or ("未取得：" + v["_error"] if "_error" in v else "—")])
    B.append(_table(["内容", "追溯号"], trace_rows, [False, False]))
    return B


# --------------------------------------------------------------------------- #
# Word 渲染
def render_docx(blocks: List[Block], out: Path) -> Path:
    from docx import Document
    from docx.enum.section import WD_ORIENT  # noqa: F401  （保留：宽表需要时可切横向）
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    def east_asia(font_el, name: str) -> None:
        rpr = font_el.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            rfonts.attrib.pop(qn(attr), None)
        rfonts.set(qn("w:eastAsia"), name)

    def shade(cell, fill: str) -> None:
        tcpr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        tcpr.append(shd)

    def borders(tbl, color: str = "DCE1E8") -> None:
        el = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement(f"w:{edge}")
            for k, v in (("val", "single"), ("sz", "4"), ("space", "0"), ("color", color)):
                e.set(qn(f"w:{k}"), v)
            el.append(e)
        tbl._tbl.tblPr.append(el)

    def write(cell, text: str, *, bold: bool = False, right: bool = False, size: float = 9, color: str = None) -> None:
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(1.5)
        if right:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = p.add_run(str(text))
        r.font.size = Pt(size)
        r.bold = bold
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin = sec.bottom_margin = Cm(2.2)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "宋体", Pt(10.5)
    east_asia(normal.element, "宋体")
    normal.paragraph_format.line_spacing = 1.35
    for name, size in (("Heading 1", 14), ("Heading 2", 11.5)):
        st_ = doc.styles[name]
        st_.font.name, st_.font.size, st_.font.bold = "黑体", Pt(size), True
        st_.font.color.rgb = RGBColor.from_string("1A2332")
        east_asia(st_.element, "黑体")
        st_.paragraph_format.space_before, st_.paragraph_format.space_after = Pt(14 if name == "Heading 1" else 8), Pt(6)

    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run("SEC 储量预评估报告（初稿）｜内部资料")
    hr.font.size, hr.font.color.rgb = Pt(8.5), RGBColor.from_string("667185")
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        run = fp.add_run()
        if kind:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            run._r.append(fc)
        else:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = text
            run._r.append(it)
        run.font.size = Pt(9)

    fig_no = 0
    for kind, v in blocks:
        if kind == "title":
            for _ in range(3):
                doc.add_paragraph()
            k = doc.add_paragraph()
            k.alignment = WD_ALIGN_PARAGRAPH.CENTER
            kr = k.add_run("SEC 单元储量预评估")
            kr.font.size, kr.font.color.rgb = Pt(11), RGBColor.from_string("1F5BB8")
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(v["title"])
            r.font.size, r.bold, r.font.name = Pt(22), True, "黑体"
            east_asia(r._element, "黑体")
            s = doc.add_paragraph()
            s.alignment = WD_ALIGN_PARAGRAPH.CENTER
            sr = s.add_run(v["subtitle"])
            sr.font.size, sr.font.color.rgb = Pt(11), RGBColor.from_string("667185")
            doc.add_paragraph()
        elif kind == "meta":
            tbl = doc.add_table(rows=0, cols=2)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            borders(tbl)
            for key, val in v:
                cells = tbl.add_row().cells
                write(cells[0], key, color="667185", size=9.5)
                shade(cells[0], "F7F9FB")
                write(cells[1], val, size=9.5)
                cells[0].width, cells[1].width = Cm(3.8), Cm(12.8)
            doc.add_paragraph()
        elif kind == "disclaimer":
            tbl = doc.add_table(rows=1, cols=1)
            borders(tbl, "E8C77A")
            write(tbl.rows[0].cells[0], "预评估声明：" + v, bold=True, size=9.5, color="7A5200")
            shade(tbl.rows[0].cells[0], "FBF3E0")
        elif kind == "pagebreak":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif kind == "h1":
            doc.add_heading(v, level=1)
        elif kind == "h2":
            doc.add_heading(v, level=2)
        elif kind == "p":
            doc.add_paragraph(v)
        elif kind == "bullets":
            for b in v:
                doc.add_paragraph(b, style="List Bullet")
        elif kind == "note":
            p = doc.add_paragraph()
            r = p.add_run("口径：" + v if not v.startswith("图件") else v)
            r.font.size, r.font.color.rgb = Pt(9), RGBColor.from_string("667185")
        elif kind == "table":
            tbl = doc.add_table(rows=1, cols=len(v["cols"]))
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            borders(tbl)
            for i, h in enumerate(v["cols"]):
                c = tbl.rows[0].cells[i]
                write(c, h, bold=True, right=v["num"][i], size=8.5, color="3F4A5A")
                shade(c, "F2F4F7")
            for ri, row in enumerate(v["rows"]):
                is_total = v["total"] and ri == len(v["rows"]) - 1
                cells = tbl.add_row().cells
                for i, val in enumerate(row):
                    write(cells[i], val, right=v["num"][i], bold=is_total, size=8.5)
                    if is_total:
                        shade(cells[i], "F7F9FB")
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        elif kind == "figure":
            fig_no += 1
            doc.add_picture(io.BytesIO(v["png"]), width=Cm(16.4))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap = doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cr = cap.add_run(f"图 {fig_no}　{v['caption']}")
            cr.font.size, cr.font.color.rgb = Pt(9), RGBColor.from_string("667185")
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    return out


# --------------------------------------------------------------------------- #
# HTML 渲染（打印即 PDF）
HTML_CSS = """
:root{--ink:#1A2332;--ink2:#3F4A5A;--muted:#667185;--rule:#DCE1E8;--soft:#EBEEF2;--head:#F2F4F7;--brand:#1F5BB8}
*{box-sizing:border-box}
body{margin:0;background:#E9EDF2;color:var(--ink);font:13px/1.7 -apple-system,BlinkMacSystemFont,"PingFang SC","Hiragino Sans GB","Microsoft YaHei","Noto Sans SC",sans-serif}
.toolbar{position:sticky;top:0;z-index:5;display:flex;gap:12px;align-items:center;justify-content:space-between;
  padding:10px 20px;background:#101B2D;color:#C5CEDB;font-size:12.5px}
.toolbar b{color:#fff;font-weight:600}
.toolbar button{height:32px;padding:0 16px;border:0;border-radius:4px;background:var(--brand);color:#fff;font:inherit;cursor:pointer}
.toolbar button:hover{filter:brightness(1.1)}
.doc{max-width:210mm;margin:24px auto 48px;background:#fff;padding:18mm 17mm;box-shadow:0 2px 18px rgba(16,27,45,.12)}
.cover{text-align:center;padding:40mm 0 12mm}
.cover .kicker{color:var(--brand);letter-spacing:.12em;font-size:13px}
.cover h1{font-size:26px;margin:10px 0 6px;letter-spacing:.02em}
.cover .sub{color:var(--muted);margin:0}
table{border-collapse:collapse;width:100%;font-size:12px;font-variant-numeric:tabular-nums}
.meta{margin:10mm 0 6mm}
.meta th{width:34%;text-align:left;color:var(--muted);font-weight:400;background:#F7F9FB}
.meta th,.meta td{border:1px solid var(--rule);padding:6px 10px}
.disc{border:1px solid #E8C77A;background:#FBF3E0;color:#7A5200;padding:10px 14px;border-radius:4px;font-weight:600}
h2{font-size:17px;margin:26px 0 10px;padding-bottom:6px;border-bottom:2px solid var(--ink)}
h3{font-size:14px;margin:18px 0 8px}
p{margin:6px 0}
ul{margin:6px 0;padding-left:20px}
li{margin:3px 0}
.tw{overflow-x:auto;margin:8px 0 12px}
.dt th{background:var(--head);color:var(--ink2);font-weight:600;text-align:left;white-space:nowrap}
.dt th,.dt td{border:1px solid var(--rule);padding:5px 8px;vertical-align:top;white-space:nowrap}
.dt td.wrap{white-space:normal;min-width:14em}
.dt .num{text-align:right;white-space:nowrap}
.dt tr.total td{font-weight:600;background:#F7F9FB}
figure{margin:12px 0 16px;text-align:center}
figure svg{width:100%;height:auto}
figcaption{color:var(--muted);font-size:12px;margin-top:4px}
.note{color:var(--muted);font-size:12px}
.pb{height:0}
@media (max-width:760px){.doc{margin:0;padding:18px 16px}.cover{padding:18mm 0 8mm}}
@page{size:A4;margin:16mm 15mm}
@media print{
  body{background:#fff}
  .toolbar{display:none}
  .doc{margin:0;padding:0;box-shadow:none;max-width:none}
  .pb{break-after:page}
  h2,h3{break-after:avoid}
  figure,tr{break-inside:avoid}
  .tw{overflow:visible}
}
"""


def render_html(blocks: List[Block]) -> str:
    e = html.escape
    title = next((v["title"] for k, v in blocks if k == "title"), "SEC 储量预评估报告")
    out: List[str] = []
    fig_no = 0
    for kind, v in blocks:
        if kind == "title":
            out.append(f'<header class="cover"><div class="kicker">SEC 单元储量预评估</div>'
                       f'<h1>{e(v["title"])}</h1><p class="sub">{e(v["subtitle"])}</p></header>')
        elif kind == "meta":
            out.append('<table class="meta">' + "".join(f"<tr><th>{e(k)}</th><td>{e(str(x))}</td></tr>" for k, x in v)
                       + "</table>")
        elif kind == "disclaimer":
            out.append(f'<div class="disc">预评估声明：{e(v)}</div>')
        elif kind == "pagebreak":
            out.append('<div class="pb"></div>')
        elif kind == "h1":
            out.append(f"<h2>{e(v)}</h2>")
        elif kind == "h2":
            out.append(f"<h3>{e(v)}</h3>")
        elif kind == "p":
            out.append(f"<p>{e(v)}</p>")
        elif kind == "bullets":
            out.append("<ul>" + "".join(f"<li>{e(b)}</li>" for b in v) + "</ul>")
        elif kind == "note":
            out.append(f'<p class="note">{e("口径：" + v if not v.startswith("图件") else v)}</p>')
        elif kind == "table":
            head = "".join(f'<th class="{"num" if v["num"][i] else ""}">{e(h)}</th>' for i, h in enumerate(v["cols"]))
            body = []
            for ri, row in enumerate(v["rows"]):
                cls = ' class="total"' if v["total"] and ri == len(v["rows"]) - 1 else ""
                body.append(f"<tr{cls}>" + "".join(
                    f'<td class="{"num" if v["num"][i] else ""}{" wrap" if i in v["wrap"] else ""}">{e(str(x))}</td>'
                    for i, x in enumerate(row)) + "</tr>")
            out.append(f'<div class="tw"><table class="dt"><thead><tr>{head}</tr></thead><tbody>{"".join(body)}</tbody></table></div>')
        elif kind == "figure":
            fig_no += 1
            out.append(f'<figure>{v["svg"]}<figcaption>图 {fig_no}　{e(v["caption"])}</figcaption></figure>')
    return ("<!doctype html><html lang=\"zh-CN\"><head><meta charset=\"utf-8\">"
            "<meta name=\"viewport\" content=\"width=device-width, initial-scale=1\">"
            f"<title>{e(title)}</title><style>{HTML_CSS}</style></head><body>"
            f'<div class="toolbar"><span><b>{e(title)}</b>　打印时选择"另存为 PDF"即可得到 PDF 版</span>'
            '<button type="button" onclick="window.print()">打印 / 另存为 PDF</button></div>'
            f'<main class="doc">{"".join(out)}</main></body></html>')


# --------------------------------------------------------------------------- #
def _safe(name: str) -> str:
    return re.sub(r"[^\w.\-]+", "_", name).strip("_")


def build(scope: str, as_of: Optional[str] = None, scenario: str = "sec",
          formats: Sequence[str] = ("docx", "html"), out_dir: Optional[Path] = None) -> Dict[str, str]:
    """生成报告文件，返回 {格式: 路径, trace_id}。评估对象或基准日不合法抛 KernelError。"""
    bad = [f for f in formats if f not in ("docx", "html")]
    if bad:
        raise S.KernelError(f"不支持的报告格式 {bad}，可选 docx / html")
    tid = trace.new_trace_id("rpt")
    data = gather(scope, as_of, scenario)
    blocks = build_blocks(data, tid)
    folder = Path(out_dir or path("artifacts_dir"))
    stem = _safe(f"SEC储量预评估报告_{data['comp']['scope']['name']}_{data['as_of']}_{scenario}")
    res: Dict[str, str] = {}
    if "docx" in formats:
        res["docx"] = str(render_docx(blocks, folder / f"{stem}.docx"))
    if "html" in formats:
        p = folder / f"{stem}.html"
        p.write_text(render_html(blocks), encoding="utf-8")
        res["html"] = str(p)
    trace.audit(tid, "kernel", "build_unit_report",
                dict(scope=scope, as_of=data["as_of"], scenario=scenario, files=list(res.values())))
    res["trace_id"] = tid
    return res
